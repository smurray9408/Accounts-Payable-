#!/usr/bin/env python3
"""
Moore Supply line-item audit.

Compares a QuickBooks "Purchase Order Detail" export (item-level, with Moore
item codes) against the line items parsed from Moore Supply invoice PDFs
(see parse_moore_invoices.py). The join key is the Moore ITEM CODE, which
appears in the QB report's `Item` column as "<code> (<description>)" and as the
ITEM NUMBER on each invoice line.

The QB report supplied is an OPEN purchase-order report (every line still
backordered), so the audit is price-centric rather than a PO-by-PO match:

  1. PO Price Check  - for each open-PO line, the unit price on the PO vs the
                       most recent price Moore actually billed for that item.
                       A material gap means the PO will (re)bill at a price that
                       disagrees with Moore's own recent invoicing.
  2. Price Anomalies - items Moore billed at INCONSISTENT unit prices across
                       invoices in the period (likely over/undercharges).
  3. No Price Ref    - open-PO items never seen on a captured invoice (cannot
                       be verified from available data).
  4. Data Coverage   - scope + the invoices whose PDF text was truncated.

Materiality: a difference is flagged when it is at least $0.02 per unit AND at
least 1% of the price. Sub-cent rounding (QB stores 3 dp, invoices print 3 dp)
is not a discrepancy.

Usage:
    python3 scripts/line_item_audit.py QB_PODETAIL.csv invoices.jsonl OUT.xlsx [DISCREPANCIES.csv]
"""
from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ABS_TOL = 0.02   # dollars per unit
REL_TOL = 0.01   # 1%


def material(a: float, b: float) -> bool:
    """True if a and b differ materially (beyond rounding)."""
    d = abs(a - b)
    base = max(abs(a), abs(b), 1e-9)
    return d >= ABS_TOL and d / base >= REL_TOL


def date_key(mmddyy: str):
    m, d, y = mmddyy.split("/")
    return (y, m, d)


def load_qb(path):
    rows = []
    with open(path) as f:
        r = csv.reader(f)
        next(r)
        for row in r:
            if len(row) < 10 or not row[2].strip():
                continue
            item = row[4]
            code = item.split()[0] if item else ""
            desc = item
            m = re.match(r"\s*\S+\s*\((.*?)\)?\s*$", item)
            if m:
                desc = m.group(1)
            rows.append(dict(
                po=row[2].strip(), date=row[1], code=code, desc=desc,
                qty=float(row[5] or 0), rcvd=float(row[6] or 0),
                bo=float(row[7] or 0), price=float(row[8] or 0),
                amount=float(row[9] or 0),
            ))
    return rows


def load_invoice_obs(path):
    """item code -> list of price observations (sales + credits)."""
    obs = defaultdict(list)
    invoices = [json.loads(l) for l in open(path)]
    for inv in invoices:
        for li in inv["line_items"]:
            obs[li["item_number"]].append(dict(
                date=inv["invoice_date"], price=li["unit_price"],
                qty=li["qty_shipped"], unit=li.get("unit"),
                invoice=inv["invoice_number"], doc_type=inv["doc_type"],
                po=inv["customer_po"],
            ))
    truncated = [inv for inv in invoices if inv["subtotal"] is None]
    return obs, invoices, truncated


# ---------- styling helpers ----------
HEAD = PatternFill("solid", fgColor="1F4E78")
HEADF = Font(bold=True, color="FFFFFF")
FLAG = PatternFill("solid", fgColor="FFC7CE")
OK = PatternFill("solid", fgColor="C6EFCE")
WARN = PatternFill("solid", fgColor="FFEB9C")
TITLE = Font(bold=True, size=14)
BOLD = Font(bold=True)


def write_header(ws, headers, row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HEAD
        cell.font = HEADF
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def latest_sale(obs_list):
    sales = [o for o in obs_list if o["doc_type"] != "CREDIT MEMO"]
    pool = sales or obs_list
    return sorted(pool, key=lambda o: date_key(o["date"]))[-1]


def build(qb_path, inv_path, out_xlsx, out_csv=None, desc_map_path=None):
    qb = load_qb(qb_path)
    obs, invoices, truncated = load_invoice_obs(inv_path)
    item_desc = {}
    if desc_map_path:
        try:
            item_desc = json.load(open(desc_map_path))
        except OSError:
            pass

    wb = Workbook()

    # ---------------- PO Price Check ----------------
    ws = wb.active
    ws.title = "PO Price Check"
    heads = ["PO #", "PO Date", "Item Code", "Description", "PO Qty",
             "PO Unit Price", "Moore Latest Price", "Latest Invoice",
             "Latest Date", "Unit Diff", "Ext. Diff (Qty x)", "Status"]
    write_header(ws, heads)
    autofit(ws, [10, 11, 11, 52, 8, 13, 15, 16, 11, 10, 15, 22])
    r = 2
    po_flags = 0
    disc_rows = []
    for x in sorted(qb, key=lambda z: (z["po"], z["code"])):
        if x["code"] in obs:
            ls = latest_sale(obs[x["code"]])
            ip = ls["price"]
            diff = ip - x["price"]
            ext = diff * x["qty"]
            if material(x["price"], ip):
                status, fill = "PRICE MISMATCH", FLAG
                po_flags += 1
                disc_rows.append(["PO Price Mismatch", x["po"], x["code"], x["desc"],
                                  x["qty"], x["price"], ip, round(ext, 2),
                                  ls["invoice"], ls["date"]])
            else:
                status, fill = "OK (within rounding)", OK
            vals = [x["po"], x["date"], x["code"], x["desc"], x["qty"],
                    round(x["price"], 3), round(ip, 3), ls["invoice"], ls["date"],
                    round(diff, 3), round(ext, 2), status]
        else:
            vals = [x["po"], x["date"], x["code"], x["desc"], x["qty"],
                    round(x["price"], 3), None, None, None, None, None,
                    "NO INVOICE PRICE REF"]
            fill = WARN
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 12:
                cell.fill = fill
        r += 1

    # ---------------- Price Anomalies ----------------
    wa = wb.create_sheet("Price Anomalies")
    heads = ["Item Code", "Description", "Distinct Sale Prices", "Min", "Max",
             "Spread $", "Spread %", "Most-Common Price", "# Invoices",
             "Detail (date / invoice / price / doc)"]
    write_header(wa, heads)
    autofit(wa, [11, 46, 16, 10, 10, 10, 9, 16, 10, 80])
    r = 2
    anomalies = 0
    code_desc = dict(item_desc)
    code_desc.update({x["code"]: x["desc"] for x in qb})
    for code, lst in sorted(obs.items()):
        sales = [o for o in lst if o["doc_type"] != "CREDIT MEMO" and o["price"] > 0]
        prices = [round(o["price"], 3) for o in sales]
        if len(set(prices)) < 2:
            continue
        mn, mx = min(prices), max(prices)
        if not material(mn, mx):
            continue
        anomalies += 1
        # most common price
        cnt = defaultdict(int)
        for p in prices:
            cnt[p] += 1
        common = max(cnt, key=lambda p: cnt[p])
        detail = "  |  ".join(
            f"{o['date']} {o['invoice']} ${o['price']:.3f} {'(CR)' if o['doc_type']=='CREDIT MEMO' else ''}"
            for o in sorted(lst, key=lambda o: date_key(o["date"]))
        )
        vals = [code, code_desc.get(code, ""), len(set(prices)), mn, mx,
                round(mx - mn, 3), f"{100*(mx-mn)/mn:.0f}%", common, len(sales), detail]
        for c, v in enumerate(vals, 1):
            wa.cell(row=r, column=c, value=v)
        wa.cell(row=r, column=1).fill = FLAG
        # add to discrepancies export
        disc_rows.append(["Inconsistent Moore Price", "", code, code_desc.get(code, ""),
                          "", mn, mx, round(mx - mn, 3), f"{cnt[mx]}x@{mx}", ""])
        r += 1

    # ---------------- No Price Reference ----------------
    wn = wb.create_sheet("No Price Ref")
    write_header(wn, ["PO #", "PO Date", "Item Code", "Description", "PO Qty", "PO Unit Price"])
    autofit(wn, [10, 11, 11, 60, 8, 13])
    r = 2
    noref = 0
    for x in sorted(qb, key=lambda z: (z["po"], z["code"])):
        if x["code"] not in obs:
            noref += 1
            for c, v in enumerate([x["po"], x["date"], x["code"], x["desc"], x["qty"], round(x["price"], 3)], 1):
                wn.cell(row=r, column=c, value=v)
            r += 1

    # ---------------- Data Coverage ----------------
    wc = wb.create_sheet("Data Coverage", 0)
    wc["A1"] = "Moore Supply — Line-Item Audit"
    wc["A1"].font = TITLE
    qb_codes = {x["code"] for x in qb}
    covered = qb_codes & set(obs)
    lines_cov = sum(1 for x in qb if x["code"] in obs)
    rows = [
        ("", ""),
        ("QuickBooks report", "Open Purchase Orders — item detail (every line still backordered)"),
        ("QB PO lines", len(qb)),
        ("QB purchase orders", len({x["po"] for x in qb})),
        ("QB distinct item codes", len(qb_codes)),
        ("", ""),
        ("Invoice source", "Moore Supply (Billtrust) PDFs, invoice dates 05/28/26 – 06/25/26"),
        ("Invoices parsed", len(invoices)),
        ("Invoice line items parsed", sum(len(i["line_items"]) for i in invoices)),
        ("Distinct item codes on invoices", len(set(obs))),
        ("", ""),
        ("Item codes with a price reference", f"{len(covered)} of {len(qb_codes)} ({100*len(covered)//max(len(qb_codes),1)}%)"),
        ("QB lines with a price reference", f"{lines_cov} of {len(qb)} ({100*lines_cov//max(len(qb),1)}%)"),
        ("", ""),
        ("FINDINGS", ""),
        ("PO lines with a material price mismatch vs Moore", po_flags),
        ("Items billed at inconsistent prices (anomalies)", anomalies),
        ("Open-PO lines with no invoice price reference", noref),
        ("", ""),
        ("Materiality threshold", "flag if difference >= $0.02/unit AND >= 1% of price"),
        ("", ""),
        ("LIMITATIONS", ""),
        ("Invoice PDF text truncated at 100k chars/file", f"{len(truncated)} tail invoices incomplete (see 'Truncated' sheet)"),
        ("Invoices for POs dated after 06/25/26", "not available (open POs not yet invoiced; email PDF download egress-blocked)"),
    ]
    rr = 2
    for k, v in rows:
        wc.cell(row=rr, column=1, value=k).font = BOLD if k in ("FINDINGS", "LIMITATIONS") else Font()
        wc.cell(row=rr, column=2, value=v)
        rr += 1
    autofit(wc, [42, 78])

    # ---------------- Truncated invoices ----------------
    wt = wb.create_sheet("Truncated")
    write_header(wt, ["Source PDF (date)", "Invoice #", "Customer PO", "Lines Captured", "Note"])
    autofit(wt, [18, 18, 14, 15, 45])
    for i, inv in enumerate(sorted(truncated, key=lambda z: z["source_file"]), 2):
        for c, v in enumerate([inv["source_file"], inv["invoice_number"], inv["customer_po"],
                               len(inv["line_items"]), "PDF text cut off before invoice total"], 1):
            wt.cell(row=i, column=c, value=v)

    wb.save(out_xlsx)

    if out_csv:
        with open(out_csv, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["Discrepancy Type", "PO #", "Item Code", "Description",
                        "Qty", "Price A", "Price B", "Diff", "Ref", "Date"])
            w.writerows(disc_rows)

    print(f"wrote {out_xlsx}")
    print(f"  PO price mismatches: {po_flags}")
    print(f"  price anomalies:     {anomalies}")
    print(f"  no price reference:  {noref}")
    print(f"  truncated invoices:  {len(truncated)}")
    return dict(po_flags=po_flags, anomalies=anomalies, noref=noref, truncated=len(truncated))


if __name__ == "__main__":
    a = sys.argv
    build(a[1], a[2], a[3],
          a[4] if len(a) > 4 else None,
          a[5] if len(a) > 5 else None)
